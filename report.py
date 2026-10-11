# -*- coding: utf-8 -*-
"""지도 결과 통지문 — 센터 「순회방문 결과지」 서식 그대로.
AI 초안(Gemini) 또는 서식 초안 → 워드(docx: 결과지 + 영양사의 당부 + 붙임 점검표) · 엑셀."""
import io
import datetime as dt

import checklist as CK

CENTER = "부천시 어린이·사회복지급식관리지원센터"

# ── 영양사의 당부 (협의사항) 라이브러리: 센터에서 실제 사용한 문안 ──
ADVICE = {
    "식품 위생·안전 관리 (식중독 예방 6대 수칙·질식 위험 식품)": """[식품 위생·안전 관리]

1. 식중독이란?
식품의 섭취로 인체에 유해한 미생물 또는 유독물질에 의해 발생하였거나 발생한 것으로 판단된 감염성 또는 독소형 질환

2. 식중독 예방 6대 수칙
 ① 손 씻기: 흐르는 물에 30초 이상 비누로 손 씻기(손바닥·손등 문지르기, 손 깍지끼고 문지르기, 손가락 마주잡고 문지르기, 엄지 손가락 돌리기, 손바닥에 손톱 밑을 문지르기)
 ② 보관온도 지키기: 냉장식품 5℃이하, 냉동식품 -18℃이하
 ③ 구분 사용하기: 날음식과 조리음식 구분하기, 칼·도마 구분 사용하기(육류용, 어패류용, 채소용, 완제품용(권장))
 ④ 가열하기: 육류·어패류 중심온도 1분 이상 익히기(육류 75℃, 어패류 85℃)
 ⑤ 세척·소독하기: 식재료와 식기류는 깨끗이 세척·소독하기(생채소·과일 염소소독: 생으로 섭취하는 채소·과일의 경우 염소소독제 100ppm에 5분 침지 후 3회 이상 세척, 식기류 소독: 열탕소독, 자외선 소독, 화학 소독 등의 방법으로 소독)
 ⑥ 끓여먹기: 물은 끓여서 마시고, 정수기 주기적으로 점검·관리하기

3. 식품 안전하게 조리하기
 ① 떡류
  - 만 1~2세) 떡류 제외, 운영위원회와 전체 학부모를 통해 떡 제공에 대한 협의가 된 시설에 한하여 센터 감수 후 식단 제공 가능 *증빙서류 확인 필요
  - 만 3~5세) 떡국 메뉴인 경우 떡이 너무 크지 않도록 제공, 기도를 막을 수 있는 위험이 있는 떡류(조랭이떡, 인절미, 백설기 등)는 식사 시간 동안 반드시 선생님들의 모니터링 필요 *2~4등분으로 잘게 잘라서 제공
 ② 매끄럽고 동글동글한 음식(방울토마토, 포도, 메추리알 등)
  - 만 1~2세) 제공 시 세로 방향으로 작게 썰고, 가급적 껍질을 제거하여 제공
  - 만 3~5세) 원형의 형태로 제공하지 않고 잘라서 제공하며 제공 시 주의
 ③ 고구마, 강력분으로 만든 빵: 퍽퍽해서 삼키기가 어렵고, 삼키면서 목이 메는 경우가 있으므로 음료와 함께 제공
 * 제공 시 주의가 필요하며, 섭취 시 모니터링 필요""",
    "어린이 급식소 염도 관리": """[어린이 급식소 염도 관리]

1. 어린이 급식소에서 제공되는 국의 권장 염도 수준은 아래와 같습니다.
 - 어린이집 및 유치원: 0.5% 이하
 - 아동 및 청소년 시설: 0.5~0.7%

2. 제공되는 국의 염도 측정 결과 (국명) (0.0)%로, ______ 관리되고 있습니다.""",
    "보존식 관리": """[보존식 관리]

1. 보존식은 제공한 모든 급식(간식 포함)을 매회 1인분 분량(100g 이상)씩 보존식 전용 용기에 담아 -18℃ 이하에서 144시간(6일) 이상 보관해야 합니다.
2. 보존식 기록지(채취 일시·메뉴명·채취자)는 식단 및 급식운영일지와 일치해야 하며, 보관 기간이 지난 보존식은 폐기 후 기록을 남겨 주시기 바랍니다.""",
    "식재료 표시사항·보관방법 준수": """[식재료 표시사항·보관방법 준수]

1. 급식소 내 표시사항이 없는 식재료는 무표시 식품으로 간주될 수 있으므로, 소분·개봉한 식재료는 소진 시까지 표시사항(제품명·유통기한·보관방법)을 함께 보관해 주시기 바랍니다.
2. 모든 식재료는 표시사항에 기재된 보관방법(냉장/냉동/실온)을 준수해야 하며, 임의로 냉동 보관하지 않도록 주의해 주시기 바랍니다.
3. 개봉한 식재료는 개봉일을 표기하고, 선입선출 원칙에 따라 사용해 주시기 바랍니다.""",
    "원산지 표시": """[원산지 표시]

1. 급식소에서 사용하는 농·축·수산물과 그 가공품은 「농수산물의 원산지 표시 등에 관한 법률」에 따라 식단표 또는 게시판에 원산지를 표시해야 합니다.
2. 식육·수산가공품은 사용하는 원료명을 정확히 표기해야 하므로 식재료 입고 시 표시사항을 확인하여 올바르게 작성해 주시기 바랍니다.
 * 예시: 생선가스 - 생선살(베트남산) X / 명태(베트남산) O""",
    "식품 알레르기 관리": """[식품 알레르기 관리]

1. 식품 알레르기 유발물질(19종)이 포함된 메뉴는 식단표에 표시하고, 알레르기 유아 명단과 대처방안(제거식·대체식·가정 지참식 등)을 조리실과 교실에 비치해 주시기 바랍니다.
2. 배식 전 알레르기 유아의 식판을 별도로 준비·확인하고, 담임교사와 조리원이 이중으로 확인하는 절차를 운영해 주시기 바랍니다.""",
    "직접 입력": "",
}


def _type_of(cl: dict) -> str:
    """점검유형: 체크리스트 이름으로 위생안전/영양 판정."""
    return "영양" if "영양" in cl["name"] and "위생" not in cl["name"] else "위생안전"


def template_draft(info: dict, cl: dict, summ: dict, defs: list, goods: list) -> dict:
    """API 키 없을 때 쓰는 규칙 기반 초안 (센터 문체)."""
    improve = []
    for d in defs:
        why = f"({d['basis']})" if d["basis"] else ""
        tail = f" ({d['note']})" if d["note"] else ""
        improve.append(f"{d['text']} 항목이 미흡합니다{why}. 개선하여 관리해 주시기 바랍니다.{tail}")
    kind = "위생·안전" if _type_of(cl) == "위생안전" else "영양"
    praise = ([f"급식소의 전반적인 {kind} 관리 수준이 우수합니다. 원장님과 조리사님의 노고에 감사드리며, 현 상태를 유지할 수 있도록 꾸준한 관리를 당부드립니다."]
              if not defs else [f"'{g['text']}' 항목은 잘 관리되고 있습니다. 현 상태를 유지해 주시기 바랍니다." for g in goods[:2]])
    return {"개선": improve, "칭찬": praise, "다음확인": [d["text"] for d in defs[:5]], "당부": "",
            "총평": f"{info.get('시설명','')} {cl['name']} 점검 결과 {summ['got']}/{summ['app_max']}점({summ['pct']}%)"}


def ai_draft(api_key: str, info: dict, cl: dict, summ: dict, defs: list, goods: list, tone: str = "따뜻하고 격려하는") -> dict:
    """Gemini로 센터 문체의 '이번 방문 시 지도사항'·칭찬·당부 초안. 실패 시 template_draft."""
    import json
    try:
        from google import genai
        client = genai.Client(api_key=api_key, vertexai=api_key.startswith("AQ."))
        body = {
            "시설": info.get("시설명"), "체크리스트": cl["name"], "점검유형": _type_of(cl), "점검일": info.get("점검일자"), "방문회차": info.get("방문회차"),
            "총점": f"{summ['got']}/{summ['app_max']} ({summ['pct']}%)",
            "감점문항": [{"번호": d["no"], "항목": d["text"], "점수": f"{d['got']}/{d['max']}", "근거": d["basis"], "점검자비고": d["note"]} for d in defs],
            "만점문항_우수": [{"번호": g["no"], "항목": g["text"], "점검자비고": g["note"]} for g in goods[:12]],
            "이전지도사항": info.get("이전지도사항", ""), "이전개선여부": info.get("이전개선여부", ""),
        }
        examples = ("문체 예시(센터 실제 통지문): "
                    "'급식소 내 표시사항이 없는 식재료는 무표시 식품으로 간주될 수 있으므로 식재료 소진 시까지 표시사항을 보관해 주시기 바랍니다. (치즈)' / "
                    "'보존식은 식단과 급식운영일지와 동일해야 하므로 기록과 메뉴가 일치할 수 있도록 관리해 주시기 바랍니다.' / "
                    "'식육·수산가공품 원료명 작성 시 사용하는 원료명을 정확히 표기해야 하므로 식재료 입고 시 표시사항을 확인하여 올바르게 작성해 주시기 바랍니다. *예시: 생선가스-생선살(베트남산) X / 명태(베트남산) O' / "
                    "칭찬 예시: '급식소의 전반적인 위생·안전 관리 수준이 우수합니다. 원장님과 조리사님의 노고에 감사드리며, 현 상태를 유지할 수 있도록 꾸준한 관리를 당부드립니다.'")
        prompt = (f"당신은 {CENTER}의 순회방문 담당 영양사입니다. 아래 점검 결과로 기관에 보내는 「순회방문 결과지」의 문안을 {tone} 어조의 한국어로 작성하세요. "
                  "반드시 JSON 객체 하나만 출력: {\"개선\": [str,...], \"칭찬\": [str,...], \"당부\": str, \"다음확인\": [str,...]}.\n"
                  "- 개선('이번 방문 시 지도사항'): 감점 문항마다 1문장씩, '~이므로/~해야 하므로 ~해 주시기 바랍니다.' 형식, 문장 끝에 점검자 비고의 구체 품목·수치를 괄호로 덧붙임. 번호·점수는 쓰지 말 것. 감점 큰 순.\n"
                  "- 칭찬: 1~2문장. 감점이 없으면 '전반적인 관리 수준이 우수합니다' 형식으로, 있으면 만점 문항·비고 중 구체적으로 잘한 점 1~2개.\n"
                  "- 당부('영양사의 당부(협의사항)'): 감점 문항과 가장 관련 있는 주제 하나를 골라 '[주제]' 제목 + 번호 매긴 교육 문안 5~10줄(기준 수치·법령 포함, 추측 금지). 감점이 없으면 계절·시기에 맞는 위생 또는 영양 주제.\n"
                  "- 다음확인: 다음 방문 때 확인할 사항 1~5개(항목명만).\n"
                  "과장·추측 금지, 비고에 없는 사실을 지어내지 말 것.\n" + examples + "\n\n" + json.dumps(body, ensure_ascii=False))
        r = client.models.generate_content(model="gemini-3.5-flash-lite", contents=prompt)
        t = r.text.strip()
        t = t[t.find("{"): t.rfind("}") + 1]
        d = json.loads(t)
        out = {k: [str(x) for x in d.get(k, [])] for k in ("개선", "칭찬", "다음확인")}
        out["당부"] = str(d.get("당부", ""))
        out["총평"] = f"{info.get('시설명','')} {cl['name']} 점검 결과 {summ['got']}/{summ['app_max']}점({summ['pct']}%)"
        return out
    except Exception:  # noqa: BLE001
        return template_draft(info, cl, summ, defs, goods)


def guidance_lines(draft: dict) -> list:
    """'이번 방문 시 지도사항' = 개선사항 + 칭찬 (센터 서식은 한 칸에 번호 매김)."""
    return [x for x in draft.get("개선", []) if x.strip()] + [x for x in draft.get("칭찬", []) if x.strip()]


# ── 워드 ──
def _shade(cell, hexcolor="EDEDED"):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hexcolor); tcPr.append(shd)


def _set_widths(table, widths_cm):
    """열 너비 고정 (Word·LibreOffice·한컴 모두 적용되도록 tblGrid·tblLayout·셀 너비를 함께 지정)."""
    from docx.shared import Cm
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    table.autofit = False
    tblPr = table._tbl.tblPr
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    mar = OxmlElement("w:tblCellMar")
    for side, v in (("top", 60), ("bottom", 60), ("left", 90), ("right", 90)):
        e = OxmlElement(f"w:{side}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); mar.append(e)
    tblPr.append(mar)
    grid = table._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths_cm):
        gc.set(qn("w:w"), str(int(Cm(w).twips)))
    for col, w in zip(table.columns, widths_cm):
        col.width = Cm(w)
    for row in table.rows:
        for c, w in zip(row.cells, widths_cm):
            c.width = Cm(w)


def _cell(cell, text, bold=False, size=9.5, align=None, shade=None):
    from docx.shared import Pt
    cell.text = ""
    lines = str(text).split("\n")
    p = cell.paragraphs[0]
    for i, ln in enumerate(lines):
        if i:
            p = cell.add_paragraph()
        run = p.add_run(ln); run.bold = bold; run.font.size = Pt(size)
        if align is not None:
            p.alignment = align
        p.paragraph_format.space_after = Pt(1)
    if shade:
        _shade(cell, shade)


def _numbered(lines: list) -> str:
    lines = [x.strip() for x in lines if x and x.strip()]
    if not lines:
        return ""
    # 이미 '1.'로 시작하면 그대로
    return "\n\n".join(x if x[:2].rstrip(".").isdigit() else f"{i}. {x}" for i, x in enumerate(lines, 1))


def to_docx(info: dict, cl: dict, summ: dict, draft: dict, defs: list, scores: dict, notes: dict, attach_sheet: bool = True) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.shared import Pt, Cm

    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2.0); sec.top_margin = sec.bottom_margin = Cm(1.8)
    st_ = doc.styles["Normal"]; st_.font.name = "맑은 고딕"; st_.font.size = Pt(10)
    st_.element.rPr.rFonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", "맑은 고딕")

    def H(text):
        p = doc.add_paragraph(); r = p.add_run(text); r.bold = True; r.font.size = Pt(12)
        p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(4)

    date_s = info.get("점검일자", "").replace("-", "/")
    t0, t1 = info.get("방문시작", ""), info.get("방문종료", "")
    when = f"{date_s} {t0} ~\n{date_s} {t1}" if (t0 or t1) else date_s
    visitors = info.get("방문자") or info.get("점검자", "")
    ptype = info.get("점검유형") or _type_of(cl)
    visit_n = info.get("방문회차", "")
    visit_s = f"{visit_n}차" if str(visit_n).isdigit() else str(visit_n)

    # 1) 머리 표
    H("순회방문 결과지")
    t = doc.add_table(rows=4, cols=4); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = [("구분", info.get("구분", "정규"), "방문일시", when), ("시설명", info.get("시설명", ""), None, None),
           ("방문자", visitors, "면담자", info.get("면담자", "")), ("점검유형", ptype, "방문회차", visit_s)]
    for r_, (a, b, c, d) in zip(t.rows, hdr):
        _cell(r_.cells[0], a, shade="EDEDED")
        if c is None:
            m = r_.cells[1].merge(r_.cells[3]); _cell(m, b)
        else:
            _cell(r_.cells[1], b); _cell(r_.cells[2], c, shade="EDEDED"); _cell(r_.cells[3], d)
    _set_widths(t, [3.2, 5.3, 3.2, 5.3])

    # 2) 지도결과(요약)
    H("지도결과 (요약)")
    prev_label = info.get("이전방문표기", "").strip()
    prev = info.get("이전지도사항", "").strip()
    prev_txt = (f"[{prev_label}]\n" if prev_label and prev else "") + (prev if prev else "해당 없음 (첫 방문)")
    improved = info.get("이전개선메모", "").strip() or info.get("이전개선여부", "")
    rows = [("방문자명(센터직원)", visitors), ("면담자명", info.get("면담자", "")), ("이전 방문 시\n지도사항", prev_txt),
            ("이전 방문 시\n지도사항 개선 여부", improved), ("이번 방문 시\n지도사항", _numbered(guidance_lines(draft)))]
    t2 = doc.add_table(rows=1 + len(rows), cols=2); t2.style = "Table Grid"; t2.alignment = WD_TABLE_ALIGNMENT.CENTER
    _cell(t2.rows[0].cells[0], "항목", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, shade="EDEDED")
    _cell(t2.rows[0].cells[1], "내용", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, shade="EDEDED")
    for r_, (a, b) in zip(t2.rows[1:], rows):
        _cell(r_.cells[0], a, bold=True); _cell(r_.cells[1], b)
    _set_widths(t2, [3.6, 13.4])

    # 3) 영양사의 당부
    advice = (draft.get("당부") or "").strip()
    if advice:
        doc.add_paragraph()
        t3 = doc.add_table(rows=1, cols=1); t3.style = "Table Grid"; t3.alignment = WD_TABLE_ALIGNMENT.CENTER
        _cell(t3.rows[0].cells[0], "영양사의 당부\n(협의사항)", bold=True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER)
        _set_widths(t3, [17.0])
        t4 = doc.add_table(rows=1, cols=2); t4.style = "Table Grid"; t4.alignment = WD_TABLE_ALIGNMENT.CENTER
        _cell(t4.rows[0].cells[0], "😊", size=16, align=WD_ALIGN_PARAGRAPH.CENTER)
        _cell(t4.rows[0].cells[1], advice)
        _set_widths(t4, [2.6, 14.4])

    doc.add_paragraph()
    p = doc.add_paragraph(); r = p.add_run(f"😊 {CENTER}"); r.bold = True; r.font.size = Pt(12); p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 4) 붙임 점검표 (선택)
    if attach_sheet:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        H(f"붙임. 점검표 — {cl['name']}  (총점 {summ['got']}/{summ['app_max']}점, {summ['pct']}%)")
        t5 = doc.add_table(rows=1, cols=5); t5.style = "Table Grid"; t5.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, hname in enumerate(["번호", "항목", "배점", "점수", "비고"]):
            _cell(t5.rows[0].cells[i], hname, bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER, shade="EDEDED")
        for part, sec_, it in CK.flat_items(cl):
            k = CK.key(part, it); v = scores.get(k)
            got = "N/A" if v == CK.NA else ("" if v is None else str(CK.value_of(v)))
            c = t5.add_row().cells
            _cell(c[0], it["no"], size=9, align=WD_ALIGN_PARAGRAPH.CENTER); _cell(c[1], f"[{sec_}] {it['text']}", size=9)
            _cell(c[2], str(it["scores"][0]), size=9, align=WD_ALIGN_PARAGRAPH.CENTER); _cell(c[3], got, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
            _cell(c[4], notes.get(k, ""), size=9)
        _set_widths(t5, [1.4, 9.2, 1.3, 1.3, 3.8])
        if cl.get("yn"):
            H("붙임 2. 집단급식소 필수 확인 사항")
            yn = info.get("필수확인", {})
            for it in cl["yn"]:
                doc.add_paragraph(f"{it['no']}. {it['text']} — {yn.get('YN|' + it['no'], '미확인')}")
    buf = io.BytesIO(); doc.save(buf); return buf.getvalue()


def to_xlsx(info: dict, cl: dict, summ: dict, scores: dict, notes: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook(); ws = wb.active; ws.title = "점검표"
    ws["A1"] = f"{cl['name']} 점검표 — {info.get('시설명', '')}"; ws["A1"].font = Font(bold=True, size=13)
    r = 2
    for k_, v_ in info.items():
        if v_ not in ("", None) and not isinstance(v_, dict):
            ws.cell(row=r, column=1, value=k_).font = Font(bold=True); ws.cell(row=r, column=2, value=str(v_)); r += 1
    r += 1
    hdr = ["영역", "구분", "번호", "항목", "근거", "배점", "점수", "비고"]
    for j, h in enumerate(hdr, 1):
        c = ws.cell(row=r, column=j, value=h); c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="E8F0FE")
    for part, sec, it in CK.flat_items(cl):
        r += 1; k = CK.key(part, it); v = scores.get(k)
        got = "N/A" if v == CK.NA else ("" if v is None else CK.value_of(v))
        for j, val in enumerate([part, sec, it["no"], it["text"], it["basis"], it["scores"][0], got, notes.get(k, "")], 1):
            ws.cell(row=r, column=j, value=val).alignment = Alignment(wrap_text=True, vertical="top")
    r += 1; ws.cell(row=r, column=4, value="총점").font = Font(bold=True); ws.cell(row=r, column=6, value=summ["app_max"]); ws.cell(row=r, column=7, value=summ["got"]).font = Font(bold=True)
    for j, w in enumerate([12, 10, 6, 60, 24, 6, 6, 30], 1): ws.column_dimensions[chr(64 + j)].width = w
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()
